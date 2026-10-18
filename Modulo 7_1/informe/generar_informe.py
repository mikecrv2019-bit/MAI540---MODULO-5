"""Genera informe/Informe_Tarea_7_1.docx (formato APA 7, 2 páginas como máximo).

Uso:  python informe/generar_informe.py
Las cifras de la auditoría por vivienda se leen ejecutando `python tareas.py vivienda`; no se transcriben a mano.
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
AUTOR = "Nombre del estudiante"          # completar antes de entregar
FUENTE = "Calibri"

# ---------- cifras: salida real del comando propio ----------
salida = subprocess.run([sys.executable, "tareas.py", "vivienda"], cwd=RAIZ, capture_output=True, text=True,
                        encoding="utf-8", env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"}, check=True).stdout
G = {}
for fila in re.finditer(r"^(sin costo|propia|alquilada)\s+(\d+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", salida, re.M):
    G[fila[1]] = dict(n=int(fila[2]), exa=float(fila[3]), malos=float(fila[4]), det=float(fila[5]))
assert set(G) == {"sin costo", "propia", "alquilada"}, salida


def f2(x):  # .77 (APA: sin cero inicial en valores acotados a 1)
    return f"{x:.2f}".lstrip("0")


# ---------- documento ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, lado, Inches(1))
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = FUENTE, Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
normal.paragraph_format.line_spacing = 2.0
normal.paragraph_format.space_after = normal.paragraph_format.space_before = Pt(0)


def campo(par, instr):
    for tipo, txt in (("begin", None), (None, instr), ("end", None)):
        run = par.add_run()
        el = OxmlElement("w:fldChar" if tipo else "w:instrText")
        if tipo:
            el.set(qn("w:fldCharType"), tipo)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


cab = sec.header.paragraphs[0]
cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
cab.paragraph_format.line_spacing = 1.0
campo(cab, "PAGE")


def parrafo(partes, indent=True, align=None, spacing=None, keep=False, sangria_francesa=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if sangria_francesa:
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(0.5), Inches(-0.5)
    if align is not None:
        p.alignment = align
    if spacing:
        p.paragraph_format.line_spacing = spacing
    p.paragraph_format.keep_with_next = keep
    for t in (partes if isinstance(partes, list) else [partes]):
        texto, est = t if isinstance(t, tuple) else (t, "")
        r = p.add_run(texto)
        r.bold, r.italic = "b" in est, "i" in est
        if "c" in est:                                   # código
            r.font.name = "Consolas"
            r.font.size = Pt(10)
    return p


def titulo(texto):
    return parrafo([(texto, "b")], indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)


def borde(celda, **lados):
    tcPr = celda._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for lado, grosor in lados.items():
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(grosor))
        el.set(qn("w:color"), "000000")
        b.append(el)


C = WD_ALIGN_PARAGRAPH.CENTER
parrafo([("Automatización de una aplicación híbrida de riesgo de crédito: un comando, un hook y una extensión", "b")],
        indent=False, align=C, spacing=1.0)
for linea in (AUTOR, "Atlantis University · MAI 540: Machine Learning", "Profesor: Kevin A. García Gallardo",
              "Tarea 7.1 · 18 de octubre de 2026"):
    parrafo(linea, indent=False, align=C, spacing=1.0)
parrafo("", indent=False, spacing=1.0)

parrafo(["La aplicación híbrida de la Clase 7 predice con un Random Forest la probabilidad de que un crédito sea malo, "
         "la explica con un modelo de lenguaje y verifica el texto antes de mostrarlo. Sobre el conjunto German Credit "
         "(Hofmann, 1994) se agregaron tres piezas de automatización, cada una con su evidencia en ", ("evidencia/", "c"), "."])

titulo("Comando propio: vivienda")
parrafo([("python tareas.py vivienda", "c"), " reemplaza una tarea que repetía a mano: auditar el predictor por grupo, lo que "
         "exige entrenar el modelo, separar la prueba, agrupar y acordarse de mirar el tamaño de cada grupo. Con la "
         f"semilla 42, la exactitud es {f2(G['propia']['exa'])} en los tres tipos de vivienda, pero el modelo detecta solo "
         f"{f2(G['propia']['det'])} de los créditos malos de quienes tienen vivienda propia (", ("n", "i"), f" = {G['propia']['n']}), frente a "
         f"{f2(G['alquilada']['det'])} en alquilada (", ("n", "i"), f" = {G['alquilada']['n']}) y {f2(G['sin costo']['det'])} en sin costo (",
         ("n", "i"), f" = {G['sin costo']['n']}, marcado como poco fiable por tener menos de 30 casos). Conviene tenerlo como comando "
         "porque la comparación se repite idéntica y el aviso de muestra pequeña no depende de la memoria de quien audita."])

titulo("Hook de Git")
parrafo([("hooks/pre-commit", "c"), " ejecuta dos comprobaciones antes de cada commit: ", ("pytest", "c"), " y una propia que "
         "cancela si queda algún bloque TODO sin completar. La segunda existe porque las pruebas solo importan ",
         ("predictor.py", "c"), " y ", ("verificador.py", "c"), ": al dejar un TODO en ", ("red_neuronal.py", "c"),
         " pytest siguió en verde (5 pruebas) y únicamente la segunda comprobación canceló el commit. La primera se "
         "demostró cambiando 0.30 por 0.99 en ", ("verificar_tono", "c"), ", lo que hizo fallar una prueba y cancelar el commit."])

titulo("Extensión auditoria-equidad")
parrafo(["Empaqueta el comando ", ("/auditoria-equidad:vivienda", "c"), ", que ejecuta el anterior y responde con una ficha de "
         "cuatro encabezados (hallazgo, evidencia, fiabilidad, siguiente comprobación), y el agente de solo lectura ",
         ("auditor-de-equidad", "c"), ", que contrasta una afirmación sobre un grupo con el código y los datos. Ambos "
         "difieren de los de clase, que auditaban por edad y revisaban explicaciones. ", ("verificar_extension.py", "c"),
         " no reporta ninguna línea FALTA. Al cargarla, ", ("plugin.json", "c"), " fija el nombre que prefija los comandos, ",
         ("vivienda.md", "c"), " define qué se ejecuta y qué puede usar Claude, y ", ("auditor-de-equidad.md", "c"),
         " le da al agente su rol y sus herramientas (Read, Grep, Glob)."])

titulo("Riesgo que reduce cada pieza y dónde no protege")
tab_titulo = parrafo([("Tabla 1", "b")], indent=False, spacing=1.0, keep=True)
parrafo([("Riesgo que evita cada pieza y un caso en el que no protege", "i")], indent=False, spacing=1.0, keep=True)

FILAS = [
    ("Comando vivienda",
     f"Quedarse con la exactitud global ({f2(G['propia']['exa'])} en los tres grupos) y olvidar el tamaño de cada grupo; "
     "la auditoría se repite igual cada vez.",
     f"Alguien lee solo la columna exactitud y concluye que el modelo es parejo, sin ver que detecta {f2(G['propia']['det'])} "
     "de los malos con vivienda propia. Además son 200 casos y una sola partición: ninguna cifra trae intervalo de confianza."),
    ("Hook pre-commit",
     "Commits con pruebas rotas y con TODO olvidados en archivos que ninguna prueba importa.",
     "git commit --no-verify lo salta. Vive en .git/hooks y no viaja con el repositorio (Chacon y Straub, 2014): "
     "quien clone debe correr instalar_hook.py. Y una regla sin prueba no se vigila: verificar_tono marca «alto» con 12 %, "
     "pero un texto que diga «considerable» pasa."),
    ("Extensión",
     "Comandos o agentes mal formados (manifiesto sin name, comando sin description) y auditorías improvisadas de cada persona.",
     "verificar_extension.py solo revisa la estructura: un comando con description pero con una instrucción equivocada se "
     "reporta como válido. El agente es un modelo de lenguaje y puede errar, y el hook no revisa la carpeta extension/."),
]
tabla = doc.add_table(rows=1 + len(FILAS), cols=3)
tabla.autofit = False
anchos = (Inches(1.1), Inches(2.4), Inches(3.0))
for i, enc in enumerate(("Pieza", "Riesgo que reduce", "Caso en que no protege")):
    tabla.rows[0].cells[i].text = ""
    r = tabla.rows[0].cells[i].paragraphs[0].add_run(enc)
    r.bold = True
for fila, datos in zip(tabla.rows[1:], FILAS):
    for i, txt in enumerate(datos):
        fila.cells[i].text = txt
for fila in tabla.rows:
    for i, celda in enumerate(fila.cells):
        celda.width = anchos[i]
        for p in celda.paragraphs:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.size = Pt(9.5)
for celda in tabla.rows[0].cells:
    borde(celda, top=8, bottom=8)
for celda in tabla.rows[-1].cells:
    borde(celda, bottom=8)

parrafo([("Nota. ", "i"), "Cifras de ", ("python tareas.py vivienda", "c"), " sobre 200 casos de prueba."],
        indent=False, spacing=1.0)
parrafo("", indent=False, spacing=1.0)
parrafo(["Las tres piezas reducen errores repetibles y ninguna sustituye el criterio: la auditoría muestra dónde mirar, el hook "
         "solo vigila lo que alguien ya escribió como regla, y la extensión empaqueta instrucciones cuya calidad "
         "no comprueba."])

titulo("Referencias")
parrafo(["Chacon, S., y Straub, B. (2014). ", ("Pro Git", "i"), " (2.ª ed.). Apress. https://git-scm.com/book"],
        sangria_francesa=True, indent=False)
parrafo(["Hofmann, H. (1994). ", ("Statlog (German Credit Data)", "i"), " [Conjunto de datos]. UCI Machine Learning "
         "Repository. https://doi.org/10.24432/C5NC77"], sangria_francesa=True, indent=False)

salida_docx = RAIZ / "informe" / "Informe_Tarea_7_1.docx"
doc.save(salida_docx)
print("Informe generado:", salida_docx)
