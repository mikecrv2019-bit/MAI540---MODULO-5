from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

d = Document()
sec = d.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(0.8)

st = d.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.0


def par(texto="", negrita=False, centro=False, cursiva=False, sangria=True, after=3):
    p = d.add_paragraph()
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif sangria:
        p.paragraph_format.first_line_indent = Inches(0.4)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    return p


def mixto(partes, sangria=True):
    p = d.add_paragraph()
    if sangria:
        p.paragraph_format.first_line_indent = Inches(0.4)
    for t, it in partes:
        r = p.add_run(t)
        r.italic = it
    return p


par("Permisos y valor de MCP en un servidor de análisis de datos con PCA", negrita=True, centro=True, after=0)
par("Estudiante: mikecrv2019 · MAI 540: Machine Learning · Atlantis University · Prof. Kevin A. García Gallardo · 6 de octubre de 2026",
    centro=True, after=5)

par("Permisos documentados", negrita=True, sangria=False, after=1)
mixto([
    ("El servidor (", False), ("mcp_server.py", True),
    (") corre por stdio y expone cinco piezas: las tools ", False), ("cargar_dataset", True),
    (" y ", False), ("ejecutar_pca", True), (", los resources ", False), ("data://datasets", True),
    (" y ", False), ("data://datasets/{nombre}", True), (" y el prompt ", False),
    ("interpretar_componentes", True),
    (". Todo lo que puede hacer se reduce a (a) listar los nombres de los CSV de ", False),
    ("datasets/", True),
    (" (hoy, iris y wine), (b) leer un CSV por su nombre exacto con ", False), ("pandas.read_csv", True),
    (" y (c) calcular PCA en memoria y devolver JSON (varianza explicada, acumulada y cargas). "
     "No escribe ni borra nada, no abre sockets de red ni importa librerías de red, y no lee el ", False),
    (".env", True),
    (" con la clave de API que vive en la misma carpeta, porque el servidor solo importa ", False),
    ("pca_utils", True), (".", False),
])
mixto([
    ("Un hallazgo que corregí durante la verificación: la versión entregada armaba la ruta como ", False),
    ("datasets/{nombre}.csv", True),
    (" sin validar, y comprobé que ", False), ("../secreto_fuera", True),
    (" leía un CSV fuera de ", False), ("datasets/", True),
    (". Ahora ", False), ("_ruta_dataset", True),
    (" acepta únicamente nombres que sean exactamente un CSV listado en ", False), ("datasets/", True),
    ("; con la corrección, ese mismo intento y ", False), ("..\\x", True),
    (" son rechazados (evidencia en ", False), ("evidencia/verificacion.log", True), (").", False),
])
mixto([
    ("Por qué es el mínimo: PCA solo necesita los valores numéricos de una tabla ya existente, así que basta lectura. "
     "El argumento es un nombre, no una ruta, porque el cliente solo debe identificar el dataset, nunca decidir "
     "dónde está en el disco. ", False),
    ("n_componentes", True),
    (" se acota entre 1 y el número de columnas numéricas, y el resultado se devuelve en lugar de guardarse, "
     "por lo que no hace falta ningún permiso de escritura. Con stdio el proceso hereda los permisos del usuario del "
     "sistema operativo; por eso el límite real lo impone el código (la lista blanca), no un sandbox.", False),
])

par("Qué aporta MCP frente a importar pca_utils.py", negrita=True, sangria=False, after=1)
mixto([
    ("Con un import directo, quien llama necesita el entorno de Python con scikit-learn, conoce los nombres y "
     "tipos de las funciones y puede alcanzar todo el módulo; por ejemplo, ", False), ("pca_utils.cargar_dataset", True),
    (" devuelve el DataFrame completo, algo que el servidor deliberadamente no publica. Como servidor MCP, la "
     "superficie queda reducida a cinco piezas con esquema de argumentos y descripciones legibles por el modelo "
     "(\"úsala antes de ejecutar_pca\"), de modo que Claude puede descubrir y encadenar las tools sin código "
     "de enlace escrito a mano. Además, el mismo servidor lo consumen sin cambios el Inspector, Claude Desktop o "
     "el ", False), ("chat.py", True),
    (" del curso, y la distinción tool/resource/prompt separa lo que el modelo decide ejecutar de lo que el cliente "
     "lee directamente y de la plantilla de interpretación ya redactada (Model Context Protocol, 2025).", False),
])

par("Referencia", negrita=True, sangria=False, after=1)
p = d.add_paragraph()
p.paragraph_format.left_indent = Inches(0.5)
p.paragraph_format.first_line_indent = Inches(-0.5)
p.add_run("Model Context Protocol. (2025). ")
p.add_run("Specification").italic = True
p.add_run(". https://modelcontextprotocol.io/specification")

d.save("Informe_Tarea_5_1.docx")
