"""Genera informe/Informe_Tarea_6_2.docx (formato APA 7) leyendo las cifras de evidencia/*.json y de la auditoria.

Uso:  python informe/generar_informe.py
Las cifras de tablas y texto se calculan desde los JSON y el informe de la Skill, no se transcriben a mano.
El formato (portada, estilos, tablas) sigue el generador de la Tarea 6.1.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
EV = RAIZ / "evidencia"
J = json.loads((EV / "busqueda_comparacion.json").read_text(encoding="utf-8"))
CV = json.loads((EV / "curvas.json").read_text(encoding="utf-8"))
AU = json.loads((EV / "auditoria_subgrupos.json").read_text(encoding="utf-8"))
IMP = json.loads((EV / "importancia_permutacion.json").read_text(encoding="utf-8"))
AUDIT_MD = (RAIZ / "auditorias" / "AUDIT_REPORT_german_credit_6_2.md").read_text(encoding="utf-8")
N_PRUEBAS = len(re.findall(r"^def test_", (RAIZ / "tests" / "test_modelo.py").read_text(encoding="utf-8"), re.M))
AUTOR = "German Credit 6.2"  # igual criterio que la portada de la Tarea 6.1; confirmar con el estudiante
FUENTE = "Calibri"

C, B, A, BASE = J["comparacion"], J["busqueda"], J["anidada"], J["datos"]["linea_base_siempre_bueno"]
DEF, AFI, RF, HGB = C["arbol_defecto"], C["arbol_afinado"], C["random_forest"], C["boosting_hgb"]
GAN = B["mejores_parametros"]
VEREDICTOS = dict(re.findall(r"^\| (V\d) [^|]+\| ([^|]+) \|$", AUDIT_MD.split("## Tabla de verificación")[0], re.M))
VEREDICTOS = {k: v.strip() for k, v in VEREDICTOS.items()}
assert len(VEREDICTOS) == 6, VEREDICTOS


# ---------- cifras ----------
def f3(x):  # .698 (APA: sin cero inicial en valores acotados a 1)
    return f"{x:.3f}".lstrip("0") if x < 1 else f"{x:.3f}"


def f4(x):  # .0036 (cuatro decimales, sin cero inicial)
    return f"{x:.4f}".lstrip("0")


def pp(x, d=1):
    return f"{x:+.{d}f}".replace("-", "−")


def m(mod, k):  # media
    return mod[k]["media"]


def d(mod, k):  # desviacion
    return mod[k]["desv"]


def sub(col):  # filas de un subgrupo
    return [f for f in AU["subgrupos"] if f["subgrupo"] == col]


# ---------- utilidades docx (iguales a la Tarea 6.1) ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, lado, Inches(1))
normal = doc.styles["Normal"]
normal.font.name = FUENTE
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
normal.paragraph_format.line_spacing = 2.0
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
h1 = doc.styles["Heading 1"]
h1.font.name, h1.font.size, h1.font.bold = FUENTE, Pt(11), True
h1.font.color.rgb = None
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    h1.element.rPr.rFonts.set(qn(a), FUENTE)
for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
    if h1.element.rPr.rFonts.get(qn(a)) is not None:
        del h1.element.rPr.rFonts.attrib[qn(a)]
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.line_spacing = 2.0
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(0)
h1.paragraph_format.keep_with_next = True


def add_field(par, instr):
    for tipo, txt in (("begin", None), (None, instr), ("end", None)):
        run = par.add_run()
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.line_spacing = 1.0
add_field(hp, "PAGE")


def parrafo(texto="", indent=True, align=None, bold=False, italic=False, spacing=None, keep=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if align is not None:
        p.alignment = align
    if spacing:
        p.paragraph_format.line_spacing = spacing
    p.paragraph_format.keep_with_next = keep
    for t in (texto if isinstance(texto, list) else [texto]):
        if isinstance(t, tuple):  # (texto, 'i') cursiva
            r = p.add_run(t[0])
            r.italic, r.bold = "i" in t[1], "b" in t[1]
        else:
            r = p.add_run(t)
            r.bold, r.italic = bold, italic
    return p


def borde(celda, **lados):
    tcPr = celda._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for lado, grosor in lados.items():
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(grosor))
        el.set(qn("w:color"), "000000")
        b.append(el)


def tabla(numero, titulo, encabezado, filas, nota, anchos):
    parrafo(f"Tabla {numero}", indent=False, bold=True, keep=True, spacing=1.0)
    parrafo(titulo, indent=False, italic=True, keep=True, spacing=1.0)
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezado))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, fila in enumerate([encabezado] + filas):
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.width = Inches(anchos[j])
            par = c.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.keep_with_next = True
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(val)
            r.font.size = Pt(9.5)
            if i == 0:
                borde(c, top=8, bottom=6)
            if i == len(filas):
                borde(c, bottom=8)
    n = parrafo([("Nota. ", "i"), nota], indent=False, spacing=1.0)
    for r in n.runs:
        r.font.size = Pt(9.5)
    n.paragraph_format.space_after = Pt(6)



def figura_curvas(ruta):
    """Figura 1 a tamano nativo (6.4 x 1.85 in) desde curvas.json, para que las fuentes sean legibles en el informe."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    paneles = [("Árbol por defecto (6.1, 20 var.)", CV["arbol_6_1_diagnostico"]),
               ("Random Forest por defecto (6.1, 20 var.)", CV["rf_6_1_diagnostico"]),
               ("HistGradientBoosting (6.2, 18 var.)", CV["boosting_hgb"])]
    fig, ejes = plt.subplots(1, 3, figsize=(6.4, 1.85), sharey=True)
    for ax, (titulo, r) in zip(ejes, paneles):
        t = r["tamanos"]
        for clave, etiqueta, color in (("entrenamiento", "entrenamiento", "tab:blue"), ("validacion", "validación", "tab:orange")):
            mu = r[f"{clave}_media"]
            de = r[f"{clave}_desv"]
            ax.plot(t, mu, "o-", ms=2.2, lw=1, color=color, label=etiqueta)
            ax.fill_between(t, [a - b for a, b in zip(mu, de)], [a + b for a, b in zip(mu, de)], color=color, alpha=0.15)
        ax.axhline(BASE, ls="--", color="gray", lw=0.8, label="«siempre bueno»")
        ax.set_title(f"{titulo}\nbrecha final {r['brecha_final']:.2f}", fontsize=6.5)
        ax.set_xlabel("casos de entrenamiento", fontsize=6.5)
        ax.tick_params(labelsize=6.5)
        ax.grid(alpha=0.25)
    ejes[0].set_ylabel("exactitud", fontsize=6.5)
    ejes[0].legend(fontsize=5.5, loc="center right")
    fig.tight_layout(pad=0.4)
    fig.savefig(ruta, dpi=250)
    plt.close(fig)


TITULO = "Optimización y auditoría de un modelo de aprobación de crédito en German Credit"

# ---------- portada ----------
for _ in range(3):
    parrafo("", indent=False)
parrafo(TITULO, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
parrafo("", indent=False)
for linea in (AUTOR, "Atlantis University", "MAI 540: Machine Learning", "Kevin A. García Gallardo", "13 de octubre de 2026"):
    parrafo(linea, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

# ---------- cuerpo ----------
parrafo(TITULO, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
parrafo([
    "Se optimizó y auditó el modelo de aprobación de crédito de la Tarea 6.1 con German Credit (1,000 solicitudes, "
    "20 variables, 30 % de créditos malos; Hofmann, 1994). Aprobar un crédito malo suele costar más que rechazar uno bueno, "
    "así que, junto con la exactitud, se reporta la proporción de créditos malos detectados (recall de la clase mala), "
    "la métrica prioritaria; no hay matriz de costos y no se supuso ninguna. Edad y estado civil se usaron solo para "
    "auditar, no para decidir (18 variables)."])

doc.add_heading("Método", level=1)
parrafo([
    "Todos los modelos usaron la misma validación cruzada estratificada de 5 pliegues (semilla 42) y un ", ("Pipeline", "i"),
    " con la codificación ajustada dentro de cada pliegue (Pedregosa et al., 2011); el protocolo reproduce las exactitudes de "
    f"la Tarea 6.1 ({f3(J['control_tarea_6_1']['arbol_defecto']['exactitud_media'])} y {f3(J['control_tarea_6_1']['random_forest']['exactitud_media'])}). "
    "El árbol se afinó con ", ("GridSearchCV", "i"), " sobre ", ("max_depth", "i"), " ∈ {1, 2, 3, 4, 5, 7, 10, sin límite} y ",
    ("min_samples_leaf", "i"), " ∈ {1, 2, 5, 10, 20} (40 combinaciones), eligiendo por exactitud. Como el mejor puntaje de "
    "una búsqueda es optimista (Cawley y Talbot, 2010), el número honesto sale de una validación cruzada anidada: la "
    "búsqueda se repite dentro de cada pliegue externo con solo su entrenamiento. Random Forest e ",
    ("HistGradientBoostingClassifier", "i"), " se evaluaron con parámetros por defecto."])

doc.add_heading("Búsqueda y medición honesta", level=1)
folds = A["por_pliegue_externo"]
prof = [f["parametros"]["max_depth"] for f in folds]
hojas = [f["parametros"]["min_samples_leaf"] for f in folds]
parrafo([
    f"La ganadora fue max_depth = {GAN['max_depth']} y min_samples_leaf = {GAN['min_samples_leaf']}: exactitud {f3(B['exactitud_busqueda_optimista'])} "
    f"en la búsqueda frente a {f3(m(DEF, 'exactitud'))} del árbol por defecto ({pp(B['mejora_optimista_pp'])} pp). Con validación anidada, el árbol "
    f"afinado obtuvo {f3(m(AFI, 'exactitud'))} (DE = {f3(d(AFI, 'exactitud'))}); la búsqueda sobrestimó {A['optimismo_busqueda_pp']:.1f} pp y la mejora honesta es "
    f"{pp(A['mejora_honesta_pp'])} pp. La ganadora cambió entre pliegues (profundidad {min(map(int, prof))} a {max(map(int, prof))}, hoja mínima "
    f"{min(hojas)} a {max(hojas)}): el óptimo es inestable. Además, afinar por exactitud redujo los malos detectados de "
    f"{f3(m(DEF, 'recall_bad'))} a {f3(m(AFI, 'recall_bad'))}."])

parrafo([
    f"En la Tabla 1, HistGradientBoosting y Random Forest superan la línea base en {(m(HGB, 'exactitud') - BASE) * 100:.1f} y "
    f"{(m(RF, 'exactitud') - BASE) * 100:.1f} pp; el árbol afinado, en {(m(AFI, 'exactitud') - BASE) * 100:.1f} pp. Su exactitud es equivalente "
    f"({f3(m(HGB, 'exactitud'))} y {f3(m(RF, 'exactitud'))}), pero HistGradientBoosting detecta {f3(m(HGB, 'recall_bad'))} de los malos frente a "
    f"{f3(m(RF, 'recall_bad'))}."])

filas1 = []
for nombre, mod in (("Árbol por defecto", DEF), ("Árbol afinado (anidada)", AFI), ("Random Forest", RF), ("HistGradientBoosting", HGB)):
    ts = m(mod, "tiempo_ajuste_s")
    filas1.append([nombre, f"{f3(m(mod, 'exactitud'))} ({f3(d(mod, 'exactitud'))})",
                   f"{f3(m(mod, 'recall_bad'))} ({f3(d(mod, 'recall_bad'))})", f"{ts:.2f}" + ("*" if mod is AFI else "")])
filas1.insert(2, ["Árbol afinado (búsqueda, optimista)", f3(B["exactitud_busqueda_optimista"]),
                  f3(B["recall_bad_busqueda_optimista"]), "—"])
tabla(1, "Comparación de modelos con el mismo protocolo (validación cruzada de 5 pliegues)",
      ["Modelo", "Exactitud M (DE)", "Malos detectados M (DE)", "Ajuste (s/pliegue)"], filas1,
      f"M = media; DE = desviación estándar muestral entre pliegues. Línea base «siempre bueno» = {f3(BASE)}. "
      f"* Incluye la búsqueda completa dentro de cada pliegue; ajustar una vez la configuración ganadora toma "
      f"{J['arbol_config_final_sin_busqueda']['tiempo_ajuste_s']['media']:.3f} s.",
      [2.3, 1.4, 1.6, 1.2])

doc.add_heading("Curva de aprendizaje", level=1)
dt, drf, hg = CV["arbol_6_1_diagnostico"], CV["rf_6_1_diagnostico"], CV["boosting_hgb"]
parrafo([
    "La Tarea 6.1 no incluyó curva de aprendizaje; el diagnóstico se reconstruyó con sus modelos por defecto y 20 variables, con el "
    "mismo código y la misma validación cruzada (Figura 1). La brecha entre entrenamiento y validación fue "
    f"{f3(dt['brecha_final'])} en el árbol y {f3(drf['brecha_final'])} en Random Forest; en HistGradientBoosting es {f3(hg['brecha_final'])} "
    f"({pp((hg['brecha_final'] - drf['brecha_final']) * 100)} pp frente a Random Forest; con 20 variables, {f3(CV['boosting_hgb_20_variables']['brecha_final'])}); no se redujo de forma relevante porque aprende casi "
    f"por completo el entrenamiento ({f3(hg['entrenamiento_final'])}). Solo el árbol afinado la reduce ({f3(CV['arbol_afinado']['brecha_final'])}), a costa de rendimiento. "
    f"La validación sigue subiendo ({pp(hg['pendiente_validacion_pp_por_100_casos'], 2)} pp por cada 100 casos entre "
    f"{hg['tamano_inicio_tramo']} y {hg['tamanos'][-1]}) sin llegar a una meseta: más datos probablemente ayudarían, poco y con incertidumbre (DE de 1 a 3 pp "
    f"entre pliegues)."])

figura_curvas(RAIZ / "informe" / "figura_curvas.png")
doc.add_picture(str(RAIZ / "informe" / "figura_curvas.png"), width=Inches(6.4))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].paragraph_format.line_spacing = 1.0
doc.paragraphs[-1].paragraph_format.keep_with_next = True
n = parrafo([("Figura 1. ", "b"), ("Curvas de aprendizaje: diagnóstico de la Tarea 6.1 (reconstruido) y modelo elegido. ", "i"),
             "Banda = ±1 DE entre pliegues; línea punteada = línea base «siempre bueno»."], indent=False, spacing=1.0)
for r in n.runs:
    r.font.size = Pt(9.5)
n.paragraph_format.space_after = Pt(6)

doc.add_heading("Auditoría y pruebas", level=1)
filas2 = []
for col, etiqueta in (("age_grupo", "Edad (4 grupos)"), ("personal_status", "Estado civil y sexo (4)"), ("housing", "Vivienda (3)")):
    fs = sub(col)
    peor = max(fs, key=lambda f: abs(f["dif_vs_global"]))
    filas2.append([etiqueta, f"{f3(min(f['recall_bad'] for f in fs))}–{f3(max(f['recall_bad'] for f in fs))}",
                   f"{peor['valor']} ({peor['positivos_bad']} malos)", pp(peor["dif_vs_global"], 3),
                   "supera" if peor["supera_umbral"] else "dentro"])
tabla(2, "Malos detectados por subgrupo, HistGradientBoosting (predicciones fuera de pliegue)",
      ["Variable", "Rango de recall", "Mayor desviación", "Dif. vs global", "Umbral .10"], filas2,
      f"Recall global = {f3(AU['recall_global'])}. Dif. = recall del subgrupo − global; mínimo {AU['min_positivos']} malos.",
      [1.9, 1.1, 2.0, 0.9, 0.8])
sup = [f for f in AU["subgrupos"] if f["supera_umbral"]]
assert [f["subgrupo"] for f in sup] == ["personal_status", "housing"], sup  # el texto nombra estos dos
parrafo([
    f"La Skill de auditoría del Módulo 4 (v2.0) dio: métricas {VEREDICTOS['V1']}; partición {VEREDICTOS['V2']}; fuga {VEREDICTOS['V3']}; "
    f"subgrupos {VEREDICTOS['V4']}; datos de entrada {VEREDICTOS['V5']}; variables no admisibles {VEREDICTOS['V6']}. "
    f"Superan el umbral (Tabla 2) «{sup[0]['valor']}» en estado civil ({f3(sup[0]['recall_bad'])}; {sup[0]['aciertos_bad']} de {sup[0]['positivos_bad']} malos: muestra pequeña) y "
    f"«{sup[1]['valor']}» en vivienda ({f3(sup[1]['recall_bad'])}; excede por {f4(abs(sup[1]['dif_vs_global']) - AU['umbral'])}). "
    "Partición y fuga fallan por diseño: el diagnóstico de la 6.1 usa otras variables y el modelo se eligió entre tres con la "
    f"misma validación cruzada, sin conjunto reservado. Las {N_PRUEBAS} pruebas de pytest pasan; dos exigen superar la línea base de {f3(BASE)} y una DE entre "
    "pliegues ≤ .05, umbral de 1.5 veces el error estándar de un pliegue de 200 casos (√(.70 × .30 / 200) = .032)."])

doc.add_heading("Recomendación", level=1)
top = IMP[:3]
fn = HGB["matriz_fuera_de_pliegue"]["FN"]
tp = HGB["matriz_fuera_de_pliegue"]["TP"]
parrafo([
    f"Llevaría a producción HistGradientBoosting, como apoyo a la decisión y no como decisión automática. Rendimiento: exactitud {f3(m(HGB, 'exactitud'))} "
    f"(DE = {f3(d(HGB, 'exactitud'))}, la menor) y {f3(m(HGB, 'recall_bad'))} de malos detectados, la más alta; Random Forest empata en exactitud pero detecta "
    f"{f3(m(RF, 'recall_bad'))}, y el árbol afinado es inestable (DE del recall = {f3(d(AFI, 'recall_bad'))}). Costo: {m(HGB, 'tiempo_ajuste_s'):.2f} s por ajuste "
    f"frente a {m(AFI, 'tiempo_ajuste_s'):.1f} s del árbol afinado con su búsqueda. Interpretabilidad: menor que la de un árbol; la "
    f"importancia por permutación indica que usa sobre todo {top[0]['variable']}, {top[1]['variable']} y {top[2]['variable']} (caídas de exactitud de "
    f"{f3(top[0]['caida_exactitud_media'])}, {f3(top[1]['caida_exactitud_media'])} y {f3(top[2]['caida_exactitud_media'])}), ninguna sensible. El error que más me "
    f"preocupa es aprobar un crédito malo: el modelo deja pasar {fn} de {tp + fn} ({fn / (tp + fn) * 100:.0f} %). Lo reflejé al preferirlo sobre Random Forest por "
    f"{(m(HGB, 'recall_bad') - m(RF, 'recall_bad')) * 100:.1f} pp de recall con exactitud equivalente. Reservas: la elección usó la misma validación cruzada "
    "(sesgo no medido), dos subgrupos superan el umbral y falta una matriz de costos para fijar el umbral de decisión."])

# ---------- referencias ----------
doc.add_heading("Referencias", level=1)
refs = [
    [("Cawley, G. C., y Talbot, N. L. C. (2010). On over-fitting in model selection and subsequent selection bias in performance evaluation. ", ""),
     ("Journal of Machine Learning Research", "i"), (", ", ""), ("11", "i"), (", 2079–2107.", "")],
    [("Hofmann, H. (1994). ", ""), ("Statlog (German Credit Data)", "i"),
     (" [Conjunto de datos]. UCI Machine Learning Repository. https://doi.org/10.24432/C5NC77", "")],
    [("Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., "
      "Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., y Duchesnay, E. "
      "(2011). Scikit-learn: Machine learning in Python. ", ""),
     ("Journal of Machine Learning Research", "i"), (", ", ""), ("12", "i"), (", 2825–2830.", "")],
]
for partes in refs:
    p = parrafo([(t, f) for t, f in partes], indent=False, spacing=1.0)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(2)

cp = doc.core_properties
cp.author, cp.title, cp.language, cp.comments = AUTOR, TITULO, "es", ""
cp.subject = "MAI 540: Machine Learning · Tarea 6.2"
salida = RAIZ / "informe" / "Informe_Tarea_6_2.docx"
doc.save(salida)
print("OK:", salida)
